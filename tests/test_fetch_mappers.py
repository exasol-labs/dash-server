from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from dash_server.consumption.formats import get_dataset_format
from dash_server.consumption.xlsx_format import write_xlsx
from dash_server.exasol.fetch_mappers import typed_fetch_mapper

# Column metadata and raw values exactly as pyexasol 3.0 reports them for a live Exasol
# (no fetch_mapper): `statement.columns()` and the unmapped row.
REAL_COLUMNS = {
    "AMOUNT": {"type": "DECIMAL", "precision": 12, "scale": 2},
    "QTY": {"type": "DECIMAL", "precision": 18, "scale": 0},
    "BIG_ID": {"type": "DECIMAL", "precision": 36, "scale": 0},
    "DAY": {"type": "DATE", "size": 4},
    "TS": {"type": "TIMESTAMP", "withLocalTimeZone": False, "size": 16},
    "TS_LTZ": {"type": "TIMESTAMP WITH LOCAL TIME ZONE", "withLocalTimeZone": True, "size": 16},
    "SPAN": {"type": "INTERVAL DAY TO SECOND", "size": 16, "precision": 2, "fraction": 3},
    "MONTHS": {"type": "INTERVAL YEAR TO MONTH", "size": 6, "precision": 2},
    "RATIO": {"type": "DOUBLE"},
    "FLAG": {"type": "BOOLEAN"},
    "NOTE": {"type": "CHAR", "size": 1, "characterSet": "ASCII"},
}
REAL_ROW = [
    "1234.56",
    42,
    "12345678901234567890",
    "2026-07-01",
    "2026-07-01 12:34:56.789000",
    "2026-07-01 12:34:56.000000",
    "+02 03:04:05.500",
    "+01-06",
    0.125,
    True,
    "x",
]


def _mapped_row() -> list[object]:
    return [typed_fetch_mapper(value, data_type) for value, data_type in zip(REAL_ROW, REAL_COLUMNS.values(), strict=True)]


def test_typed_fetch_mapper_converts_real_pyexasol_values() -> None:
    assert _mapped_row() == [
        Decimal("1234.56"),
        42,
        12345678901234567890,
        date(2026, 7, 1),
        datetime(2026, 7, 1, 12, 34, 56, 789000),
        datetime(2026, 7, 1, 12, 34, 56),
        # pyexasol 3.0's interval parser raises on Exasol's default DAY(2) rendering.
        "+02 03:04:05.500",
        "+01-06",
        0.125,
        True,
        "x",
    ]


def test_typed_fetch_mapper_keeps_unparseable_values_as_text() -> None:
    # A session with a non-ISO NLS_DATE_FORMAT / NLS_TIMESTAMP_FORMAT.
    assert typed_fetch_mapper("01.07.2026", REAL_COLUMNS["DAY"]) == "01.07.2026"
    assert typed_fetch_mapper("01.07.2026 12:34", REAL_COLUMNS["TS"]) == "01.07.2026 12:34"
    assert typed_fetch_mapper("n/a", REAL_COLUMNS["AMOUNT"]) == "n/a"
    assert typed_fetch_mapper(None, REAL_COLUMNS["AMOUNT"]) is None


def test_only_xlsx_exports_request_typed_values() -> None:
    xlsx = get_dataset_format("xlsx")
    csv = get_dataset_format("csv")
    assert xlsx is not None and xlsx.fetch_mapper is typed_fetch_mapper
    # CSV must stay byte-for-byte Exasol's own rendering.
    assert csv is not None and csv.fetch_mapper is None


def test_xlsx_export_of_real_pyexasol_values_writes_typed_cells(tmp_path: Path) -> None:
    path = tmp_path / "export.xlsx"
    write_xlsx(
        path,
        columns=list(REAL_COLUMNS),
        batches=iter([[_mapped_row()]]),
        max_rows=10,
        max_bytes=1_000_000,
        cancelled=lambda: False,
        provenance={"output_id": "typed"},
    )
    sheet = load_workbook(path)["typed"]
    cells = {name: cell for name, cell in zip(REAL_COLUMNS, next(sheet.iter_rows(min_row=2)), strict=True)}

    assert (cells["AMOUNT"].value, cells["AMOUNT"].data_type) == (1234.56, "n")
    assert (cells["QTY"].value, cells["QTY"].data_type) == (42, "n")
    # 20 digits exceeds Excel's 15 significant digits: kept exact, as text.
    assert (cells["BIG_ID"].value, cells["BIG_ID"].data_type) == ("12345678901234567890", "s")
    assert cells["DAY"].is_date and cells["DAY"].value == datetime(2026, 7, 1)
    assert cells["TS"].is_date and cells["TS"].value == datetime(2026, 7, 1, 12, 34, 56, 789000)
    assert cells["TS_LTZ"].is_date and cells["TS_LTZ"].value == datetime(2026, 7, 1, 12, 34, 56)
    assert (cells["SPAN"].value, cells["SPAN"].data_type) == ("+02 03:04:05.500", "s")
    assert (cells["MONTHS"].value, cells["MONTHS"].data_type) == ("+01-06", "s")
    assert (cells["RATIO"].value, cells["RATIO"].data_type) == (0.125, "n")
    assert (cells["FLAG"].value, cells["FLAG"].data_type) == (True, "b")
    assert (cells["NOTE"].value, cells["NOTE"].data_type) == ("x", "s")


def test_xlsx_writes_values_excel_cannot_hold_exactly_as_text(tmp_path: Path) -> None:
    path = tmp_path / "export.xlsx"
    row = [
        10**15 - 1,  # 15 digits: still a number
        10**15,  # 16 digits: Excel would round it
        Decimal("12345678901234.5"),  # 15 significant digits
        Decimal("1234567890123.4567"),  # 17 significant digits
        Decimal("1234.5000"),  # trailing zeros aren't significant
        date(1899, 12, 31),  # before Excel's 1900 date system
    ]
    write_xlsx(
        path,
        columns=[f"C{i}" for i in range(len(row))],
        batches=iter([[row]]),
        max_rows=10,
        max_bytes=1_000_000,
        cancelled=lambda: False,
        provenance={"output_id": "limits"},
    )
    cells = [(cell.value, cell.data_type) for cell in next(load_workbook(path)["limits"].iter_rows(min_row=2))]
    assert cells == [
        (999999999999999, "n"),
        ("1000000000000000", "s"),
        (12345678901234.5, "n"),
        ("1234567890123.4567", "s"),
        (1234.5, "n"),
        ("1899-12-31", "s"),
    ]

